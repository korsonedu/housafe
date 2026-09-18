# -*- coding: utf-8 -*-
"""家安 housafe · 融资计划独立 PPTX 生成器
2 页：融资概览（深色）+ 里程碑与进展（浅色）
视觉：瑞士国际主义 + 克莱因蓝 IKB
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ================= 视觉常量 =================
IKB       = RGBColor(0x00, 0x2F, 0xA7)
IKB_BRT   = RGBColor(0x5B, 0x7B, 0xFF)
INK       = RGBColor(0x0A, 0x0A, 0x0A)
PAPER     = RGBColor(0xFA, 0xFA, 0xF8)
GREY1     = RGBColor(0xF0, 0xF0, 0xEE)
GREY2     = RGBColor(0xD4, 0xD4, 0xD2)
GREY3     = RGBColor(0x73, 0x73, 0x73)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
INK_CARD  = RGBColor(0xF5, 0xF5, 0xF4)

SANS = 'Arial'
ZH   = '微软雅黑'

EMU_W, EMU_H = Inches(13.333), Inches(7.5)
M = 0.55  # 页边距

# ================= 绘制原语 =================
def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])

def bg(slide, color):
    r = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, EMU_W, EMU_H)
    r.fill.solid(); r.fill.fore_color.rgb = color
    r.line.fill.background(); r.shadow.inherit = False
    sp = r._element; sp.getparent().remove(sp); slide.shapes._spTree.insert(2, sp)
    return r

def rect(slide, x, y, w, h, fill=None, line=None, line_w=1.0):
    r = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None: r.fill.background()
    else: r.fill.solid(); r.fill.fore_color.rgb = fill
    if line is None: r.line.fill.background()
    else: r.line.color.rgb = line; r.line.width = Pt(line_w)
    r.shadow.inherit = False
    return r

def hline(slide, x, y, w, color, weight=1.0):
    ln = slide.shapes.add_connector(2, Inches(x), Inches(y), Inches(x+w), Inches(y))
    ln.line.color.rgb = color; ln.line.width = Pt(weight)
    ln.shadow.inherit = False
    return ln

def _ea(run, ea):
    rPr = run._r.get_or_add_rPr()
    e = rPr.find(qn('a:ea'))
    if e is None:
        e = rPr.makeelement(qn('a:ea'), {}); rPr.append(e)
    e.set('typeface', ea)

def add_run(p, text, size, color=INK, bold=False, italic=False, latin=SANS, ea=ZH, spc=None):
    r = p.add_run(); r.text = text
    f = r.font; f.size = Pt(size); f.bold = bold; f.italic = italic
    f.color.rgb = color; f.name = latin
    _ea(r, ea)
    if spc is not None:
        r._r.get_or_add_rPr().set('spc', str(int(spc*100)))
    return r

def box(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tf

def put(slide, x, y, w, h, text, size, color=INK, bold=False, align=PP_ALIGN.LEFT,
        anchor=MSO_ANCHOR.TOP, latin=SANS, ea=ZH, italic=False, spc=None, lh=None):
    tf = box(slide, x, y, w, h, anchor)
    p = tf.paragraphs[0]; p.alignment = align
    if lh: p.line_spacing = lh
    add_run(p, text, size, color, bold, italic, latin, ea, spc)
    return tf

def header(slide, left, right, dark=False):
    c = WHITE if dark else GREY3
    put(slide, M, 0.36, 8, 0.3, left, 9.5, c, latin=SANS, ea=ZH, spc=1.2)
    put(slide, 6.333, 0.36, 6.45, 0.3, right, 9.5, c, align=PP_ALIGN.RIGHT, latin=SANS, spc=1.2)

def kicker(slide, x, y, text, color=IKB):
    put(slide, x, y, 9, 0.3, text, 10.5, color, bold=True, latin=SANS, ea=ZH, spc=1.5)

def title(slide, x, y, text, size=30, color=INK):
    tf = box(slide, x, y, 12.2, size/40+0.6)
    p = tf.paragraphs[0]; p.line_spacing = 1.05
    for i, line in enumerate(text.split('\n')):
        if i: p = tf.add_paragraph(); p.line_spacing = 1.05
        add_run(p, line, size, color, bold=False)
    return tf

# ================= 第 1 页：融资概览（深色） =================
def page1_funding_overview(prs):
    s = blank(prs); bg(s, INK)
    header(s, '家安 housafe · 融资计划', '天使轮 · 01 / 02', dark=True)
    kicker(s, M, 0.85, 'SEED ROUND · 500 万人民币', IKB_BRT)
    title(s, M, 1.2, '融资概览', 27, WHITE)

    # ── 第一行：融资额（大卡）+ 跑道 + 出让 ──
    row1_y = 2.1
    big_h = 2.35

    # 左：融资额 IKB 大卡
    rect(s, M, row1_y, 4.15, big_h, fill=IKB)
    put(s, M+0.28, row1_y+0.22, 3.6, 0.3, '融资金额', 9.5, WHITE, spc=1.2)
    put(s, M+0.28, row1_y+0.55, 3.6, 0.85, '500 万', 36, WHITE, latin=SANS)
    put(s, M+0.28, row1_y+1.45, 3.6, 0.7,
        '人民币 · 出让约 14–17%\n投前估值约 2,500–3,000 万',
        10, RGBColor(0xDD, 0xE4, 0xF5), lh=1.35)

    # 中上：跑道
    mid_x = M + 4.45
    rect(s, mid_x, row1_y, 4.1, 1.05, fill=GREY1)
    put(s, mid_x+0.22, row1_y+0.15, 3.7, 0.25, '融资跑道', 9, GREY3, spc=1.1)
    put(s, mid_x+0.22, row1_y+0.42, 3.7, 0.5, '12 个月', 24, INK, latin=SANS)

    # 中下：融资阶段
    rect(s, mid_x, row1_y+1.3, 4.1, 1.05, fill=GREY1)
    put(s, mid_x+0.22, row1_y+1.45, 3.7, 0.25, '融资阶段', 9, GREY3, spc=1.1)
    put(s, mid_x+0.22, row1_y+1.72, 3.7, 0.5, '天使轮', 24, INK)

    # 右上：融资后目标
    right_x = M + 8.85
    rect(s, right_x, row1_y, 3.93, big_h, fill=GREY1)
    put(s, right_x+0.22, row1_y+0.22, 3.5, 0.25, '融资后目标', 9, GREY3, spc=1.1)
    goals = [
        '1,000+ 户付费装机',
        '误报 < 2 次/周/户',
        '续费率数据初步验证',
        '启动 A 轮融资',
    ]
    for i, g in enumerate(goals):
        c = IKB if i == 1 else INK
        put(s, right_x+0.22, row1_y+0.52 + i*0.42, 3.5, 0.38, f'— {g}', 10.5, c, lh=1.2)

    # ── 第二行：资金用途详情 ──
    row2_y = 4.7
    hline(s, M, row2_y - 0.1, 12.23, RGBColor(0x44, 0x44, 0x44), 0.8)

    # 标题
    put(s, M, row2_y, 6, 0.3, '资金用途明细', 9.5, WHITE, spc=1.2)

    # 四个用途卡片横排
    uses = [
        ('研发 / 团队', '275 万', '55%', '5→8 人 · 含创始人薪资\nAI 模型训练 · 后端 · App', True),
        ('硬件 / 供应链', '75 万', '15%', '雷达开发板 · 模组采购\n小批量试产 · 工装夹具', False),
        ('获客 / 运营', '100 万', '20%', '种子用户获取 · 社区冷启\n适老化补贴渠道对接', False),
        ('云 / 数据 / 其他', '50 万', '10%', 'GPU 云 · TimescaleDB\n办公 · 行政 · 合规', False),
    ]
    uy = row2_y + 0.4
    uw = 2.92
    ug = 0.18
    for i, (label, amount, pct, detail, accent) in enumerate(uses):
        ux = M + i * (uw + ug)
        fill = IKB if accent else INK_CARD
        rect(s, ux, uy, uw, 2.0, fill=fill)

        tc = WHITE if accent else GREY3
        nc = WHITE if accent else INK
        dc = RGBColor(0xDD, 0xE4, 0xF5) if accent else RGBColor(0x44, 0x44, 0x44)

        put(s, ux+0.18, uy+0.15, uw-0.36, 0.25, label, 9, tc, spc=1.0)
        # 金额 + 百分比放同一行
        tf = box(s, ux+0.18, uy+0.42, uw-0.36, 0.55)
        p = tf.paragraphs[0]; p.line_spacing = 1.1
        add_run(p, amount, 20, nc, latin=SANS)
        add_run(p, f'  {pct}', 12, tc if accent else GREY3, latin=SANS, spc=0.8)
        put(s, ux+0.18, uy+1.05, uw-0.36, 0.85, detail, 9, dc, lh=1.3)

    # 底部注脚：已完成工作概要
    put(s, M, 7.0, 12.23, 0.4,
        '当前进展：地基系统 10 Tasks ✅ + AI 引擎 Phase A-1/2/3 ✅ — 软件骨架已跑通，本轮资金主要投向 AI 模型训练 + 硬件采购 + 商业化启动',
        8.5, RGBColor(0x99, 0x99, 0x99), lh=1.2)


# ================= 第 2 页：里程碑 + 当前进展（浅色） =================
def page2_milestones(prs):
    s = blank(prs); bg(s, PAPER)
    header(s, '家安 housafe · 融资计划', '天使轮 · 02 / 02')
    kicker(s, M, 0.85, 'ROADMAP & PROGRESS')
    title(s, M, 1.2, '关键里程碑与当前进展', 27)

    # ── 上半部分：12 个月里程碑时间线 ──
    put(s, M, 2.1, 6, 0.3, '12 个月关键里程碑', 11, INK, bold=False, spc=1.0)

    axis_y = 3.35
    x0 = 1.6
    span = 10.0
    hline(s, M, axis_y, 12.23, GREY2, 1.5)

    milestones = [
        ('M0–3', 'AI 模型上线', '跌倒 + 生命体征\n内测 5–10 户', 'up', True),
        ('M3–6', '商用版上线', '全屋套 S/L 发布\n开始公开售卖', 'down', False),
        ('M6–9', '规模化验证', '500 户装机\n误报优化到位', 'up', False),
        ('M9–12', '数据飞轮', '1,000+ 户\n启动 A 轮', 'down', True),
    ]

    for i, (period, name, detail, direction, accent) in enumerate(milestones):
        cx = x0 + i * (span / 3)

        # 节点圆点
        dot_size = 0.14
        dot_fill = IKB if accent else INK
        rect(s, cx - dot_size/2, axis_y - dot_size/2, dot_size, dot_size, fill=dot_fill)

        if direction == 'up':
            # 标签在上方
            put(s, cx - 1.4, axis_y - 1.65, 2.8, 0.25, period, 9.5, IKB if accent else GREY3,
                align=PP_ALIGN.CENTER, latin=SANS, spc=0.6)
            put(s, cx - 1.4, axis_y - 1.35, 2.8, 0.35, name, 13, IKB if accent else INK,
                align=PP_ALIGN.CENTER)
            put(s, cx - 1.45, axis_y - 0.92, 2.9, 0.55, detail, 9, GREY3,
                align=PP_ALIGN.CENTER, lh=1.25)
        else:
            # 标签在下方
            put(s, cx - 1.4, axis_y + 0.22, 2.8, 0.25, period, 9.5, IKB if accent else GREY3,
                align=PP_ALIGN.CENTER, latin=SANS, spc=0.6)
            put(s, cx - 1.4, axis_y + 0.52, 2.8, 0.35, name, 13, IKB if accent else INK,
                align=PP_ALIGN.CENTER)
            put(s, cx - 1.45, axis_y + 0.95, 2.9, 0.55, detail, 9, GREY3,
                align=PP_ALIGN.CENTER, lh=1.25)

    # ── 下半部分：当前已完成工作 ──
    hline(s, M, 4.7, 12.23, GREY2, 1)
    put(s, M, 4.85, 6, 0.3, '当前已完成工作', 11, INK, bold=False, spc=1.0)

    # 三列：地基 / AI 引擎 / 前端
    done_cols = [
        ('⚙️ 地基系统', IKB, [
            'Django + DRF 后端',
            'JWT 鉴权（注册/登录/刷新）',
            '事件契约包（Tier-1 感知事件）',
            'TimescaleDB 时序存储 + 查询',
            'WebSocket 实时通道（JWT+家庭隔离）',
            '设备绑定/列表/重命名',
            '仿真器 + 端到端冒烟测试',
        ]),
        ('🧠 AI 引擎', IKB_BRT, [
            '世界模型共享类型定义',
            'Redis Stream 消费端',
            'Typed Frame 解析',
            'FallbackEngine（3 条确定性规则）',
            'ELPD Encoder/Decoder 架构就绪',
        ]),
        ('📱 前端 App', IKB, [
            'Expo 项目框架',
            '主题系统（暖白+青绿养老基调）',
            '登录/注册流程',
            '实时 Today 首页',
            '家庭/老人管理界面',
        ]),
    ]

    col_w = 3.94
    col_g = 0.24
    col_y = 5.25

    for i, (col_title, accent_color, items) in enumerate(done_cols):
        cx = M + i * (col_w + col_g)
        # 列标题
        rect(s, cx, col_y, col_w, 0.42, fill=INK_CARD)
        put(s, cx+0.18, col_y+0.06, col_w-0.36, 0.3, col_title, 11, INK, bold=False)

        # 条目
        for j, item in enumerate(items):
            iy = col_y + 0.5 + j * 0.25
            c = GREY3 if j < len(items) else accent_color
            put(s, cx+0.18, iy, col_w-0.36, 0.23, f'✓  {item}', 9, c, lh=1.15)

    # 底部总结
    hline(s, M, 6.85, 12.23, GREY2, 0.8)
    put(s, M, 6.95, 12.23, 0.4,
        '软件骨架已跑通。本轮 500 万天使轮资金主要投向 AI 模型训练与调优、雷达硬件采购与试产、种子用户获取——从"能跑"到"能卖"。',
        9.5, RGBColor(0x40, 0x40, 0x40), lh=1.3)


# ================= 生成入口 =================
def build():
    prs = Presentation()
    prs.slide_width = EMU_W
    prs.slide_height = EMU_H

    page1_funding_overview(prs)
    page2_milestones(prs)

    out = '融资计划.pptx'
    prs.save(out)
    print(f'✅ 融资计划: {len(prs.slides)} 页 → {out}')


if __name__ == '__main__':
    build()
