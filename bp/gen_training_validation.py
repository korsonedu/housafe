# -*- coding: utf-8 -*-
"""家安 housafe · 世界模型训练验证 — 独立单页 PPTX 生成器

复用 pptx_gen 的母版层（视觉常量 + 绘制原语），生成一张 dark 主题的
「世界模型 — 小样本即能用」验证页。

输出: 世界模型_训练验证.pptx
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ========== 从 pptx_gen 复用母版（直接 copy 的常量/原语，不 import 整个模块） ==========

IKB       = RGBColor(0x00, 0x2F, 0xA7)
IKB_BRT   = RGBColor(0x5B, 0x7B, 0xFF)
INK       = RGBColor(0x0A, 0x0A, 0x0A)
PAPER     = RGBColor(0xFA, 0xFA, 0xF8)
GREY1     = RGBColor(0xF0, 0xF0, 0xEE)
GREY2     = RGBColor(0xD4, 0xD4, 0xD2)
GREY3     = RGBColor(0x73, 0x73, 0x73)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
INK_CARD  = RGBColor(0xF5, 0xF5, 0xF4)
DARK_CARD = RGBColor(0x15, 0x15, 0x15)

SANS = 'Arial'
ZH   = '微软雅黑'

EMU_W, EMU_H = Inches(13.333), Inches(7.5)
M = 0.55


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def bg(slide, color):
    r = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, EMU_W, EMU_H)
    r.fill.solid(); r.fill.fore_color.rgb = color
    r.line.fill.background(); r.shadow.inherit = False
    sp = r._element; sp.getparent().remove(sp); slide.shapes._spTree.insert(2, sp)
    return r


def rect(slide, x, y, w, h, fill=None, line=None, line_w=1.0):
    r = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                               Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None: r.fill.background()
    else: r.fill.solid(); r.fill.fore_color.rgb = fill
    if line is None: r.line.fill.background()
    else: r.line.color.rgb = line; r.line.width = Pt(line_w)
    r.shadow.inherit = False
    return r


def hline(slide, x, y, w, color, weight=1.0):
    ln = slide.shapes.add_connector(2, Inches(x), Inches(y),
                                    Inches(x + w), Inches(y))
    ln.line.color.rgb = color; ln.line.width = Pt(weight)
    ln.shadow.inherit = False
    return ln


def _ea(run, ea):
    rPr = run._r.get_or_add_rPr()
    e = rPr.find(qn('a:ea'))
    if e is None:
        e = rPr.makeelement(qn('a:ea'), {}); rPr.append(e)
    e.set('typeface', ea)


def add_run(p, text, size, color=INK, bold=False, italic=False,
            latin=SANS, ea=ZH, spc=None):
    r = p.add_run(); r.text = text
    f = r.font; f.size = Pt(size); f.bold = bold; f.italic = italic
    f.color.rgb = color; f.name = latin
    _ea(r, ea)
    if spc is not None:
        r._r.get_or_add_rPr().set('spc', str(int(spc * 100)))
    return r


def box(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tf


def put(slide, x, y, w, h, text, size, color=INK, bold=False,
        align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, latin=SANS, ea=ZH,
        italic=False, spc=None, lh=None):
    tf = box(slide, x, y, w, h, anchor)
    p = tf.paragraphs[0]; p.alignment = align
    if lh: p.line_spacing = lh
    add_run(p, text, size, color, bold, italic, latin, ea, spc)
    return tf


def multi_put(slide, x, y, w, segments, lh=None, anchor=MSO_ANCHOR.TOP):
    """多样式文本段: [(text, size, color, bold?, italic?, latin?, ea?, spc?), ...]"""
    tf = box(slide, x, y, w, 2.0, anchor)
    p = tf.paragraphs[0]
    if lh: p.line_spacing = lh
    for seg in segments:
        text = seg[0]; size = seg[1]
        color = seg[2] if len(seg) > 2 else INK
        bold = seg[3] if len(seg) > 3 else False
        italic = seg[4] if len(seg) > 4 else False
        latin = seg[5] if len(seg) > 5 else SANS
        ea = seg[6] if len(seg) > 6 else ZH
        spc = seg[7] if len(seg) > 7 else None
        add_run(p, text, size, color, bold, italic, latin, ea, spc)
    return tf


# ========== 单页内容：世界模型 · 训练验证 ==========

def build_slide(prs):
    s = blank(prs); bg(s, INK)

    # ── 页眉 ──
    put(s, M, 0.36, 8, 0.3, '家安 housafe · 世界模型训练验证',
        9.5, RGBColor(0xAA, 0xAA, 0xAA), latin=SANS, ea=ZH, spc=1.2)
    put(s, 8.5, 0.36, 4.28, 0.3, 'TRAINING VALIDATION',
        9.5, RGBColor(0xAA, 0xAA, 0xAA), align=PP_ALIGN.RIGHT, latin=SANS, spc=1.2)

    # ── Kicker + 标题 ──
    put(s, M, 0.75, 9, 0.3, 'PROOF OF ARCHITECTURE', 10, IKB_BRT,
        bold=True, latin=SANS, spc=1.5)

    tf = box(s, M, 1.1, 7.6, 0.6)
    p = tf.paragraphs[0]; p.line_spacing = 1.05
    add_run(p, '世界模型 — 小样本即能用', 24, WHITE, bold=False)

    put(s, M, 1.7, 7.6, 0.4,
        '仅需 7 人采集数据即可建立高精度人体状态模型，单用户冷启动 < 14 天',
        10, RGBColor(0xCC, 0xCC, 0xCC), lh=1.2)

    # ── 架构简化图（文字版） ──
    arch_y = 2.22
    arch_boxes = [
        (M,           '毫米波点云\n稀疏 10-50 点/帧', False),
        (M + 1.65,    'Encoder\nPointNet++\n+ TCN', True),
        (M + 3.3,     '隐状态 S_t\n256-d 向量', False),
        (M + 4.95,    'Predictor\nCausal TF\n预测 Ŝ_{t+1}', True),
        (M + 6.6,     '预测偏差\nŜ ≠ S → 异常', False),
    ]
    bw, bh = 1.42, 0.88
    for i, (ax, label, accent) in enumerate(arch_boxes):
        fill = IKB if accent else DARK_CARD
        line_c = IKB_BRT if accent else RGBColor(0x44, 0x44, 0x44)
        rect(s, ax, arch_y, bw, bh, fill=fill, line=line_c, line_w=1.2)
        tc = WHITE if accent else RGBColor(0xCC, 0xCC, 0xCC)
        put(s, ax + 0.1, arch_y + 0.1, bw - 0.2, bh - 0.2,
            label, 8, tc, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, lh=1.25)
        # 箭头
        if i < 4:
            arrow_x = ax + bw + 0.02
            arrow_y = arch_y + bh / 2
            put(s, arrow_x, arrow_y - 0.12, 0.2, 0.24, '→', 12,
                IKB_BRT, align=PP_ALIGN.CENTER, latin=SANS)

    # ── 四验证指标卡片 ──
    card_y = 3.35
    card_w = 1.76
    card_h = 2.55
    card_gap = 0.12
    h_data = [
        ('H1', '可编码', '≥ 85%', '线性分类\n12 动作准确率',
         '点云稀疏≠信息不够\n——雷达能"看懂"\n人在干什么'),
        ('H2', '连续性', '≥ 0.85', '过渡帧\n余弦相似度',
         '"站→坐"是连续渐变\n隐空间像物理空间\n一样平滑'),
        ('H3', '可预测', '30%+', '预测精度\nvs"不变"基线',
         '系统能预测老人\n下一步状态——预测\n错了才知道出事了'),
        ('H4', '异常检测', '≥ 0.90', '跌倒 vs 正常\nAUC',
         '从未见过跌倒\n但能发现"这不像\n正常行为"'),
    ]
    for i, (label, name, value, metric, meaning) in enumerate(h_data):
        cx = M + i * (card_w + card_gap)
        accent = (i == 3)
        fill = IKB if accent else DARK_CARD
        line_c = IKB_BRT if accent else RGBColor(0x44, 0x44, 0x44)
        rect(s, cx, card_y, card_w, card_h, fill=fill, line=line_c, line_w=1)

        tc_label = WHITE if accent else IKB_BRT
        put(s, cx + 0.12, card_y + 0.15, card_w - 0.24, 0.25,
            label, 9, tc_label, bold=True, latin=SANS, spc=1.0)
        put(s, cx + 0.12, card_y + 0.42, card_w - 0.24, 0.22,
            name, 9.5, WHITE if accent else RGBColor(0xDD, 0xDD, 0xDD))
        # 大数字
        put(s, cx + 0.12, card_y + 0.72, card_w - 0.24, 0.6,
            value, 24, WHITE, bold=True, latin=SANS)
        # 指标名
        put(s, cx + 0.12, card_y + 1.32, card_w - 0.24, 0.4,
            metric, 7.5, RGBColor(0xAA, 0xBB, 0xDD) if accent else RGBColor(0x99, 0x99, 0x99),
            lh=1.2)
        # 分隔线
        hline(s, cx + 0.12, card_y + 1.75, card_w - 0.24,
              RGBColor(0x55, 0x75, 0xBB) if accent else RGBColor(0x44, 0x44, 0x44), 0.5)
        # 解释
        put(s, cx + 0.12, card_y + 1.85, card_w - 0.24, 0.65,
            meaning, 7.5, RGBColor(0xC8, 0xD8, 0xF0) if accent else RGBColor(0x88, 0x88, 0x88),
            lh=1.25)

    # H4 标记
    put(s, M + 3 * (card_w + card_gap), card_y + card_h + 0.05,
        card_w, 0.2, '★ 最终判决假设', 7, IKB_BRT, align=PP_ALIGN.CENTER, latin=SANS, spc=0.8)

    # ── 底部说明条 ──
    put(s, M, 6.1, 7.8, 0.7,
        '只在正常动作上训练，从未见过跌倒——但跌倒发生时，Predictor 预测误差显著高于正常'
        '（t-test p < 0.01，effect size > 2×）。'
        '这证明了"预测偏差 = 异常"机制成立。规则引擎靠"写死跌倒长什么样"来检测，'
        '世界模型靠"发现不对劲"——后者能捕获规则没写过的危险模式。',
        8, RGBColor(0x99, 0x99, 0x99), lh=1.3)

    # ═══════════════════════════════════════════
    # 右侧 40%：数据效率对比表 + 结论
    # ═══════════════════════════════════════════

    rx = 8.6
    rw = 4.2

    # 分区标题
    hline(s, rx, 1.1, rw, RGBColor(0x55, 0x55, 0x55), 1)
    put(s, rx, 1.2, rw, 0.25, '数据效率对比', 10, WHITE, bold=True, spc=1.0)
    put(s, rx, 1.45, rw, 0.25, '家安世界模型 vs 典型大模型', 8, RGBColor(0x99, 0x99, 0x99))

    # 表格数据
    table_rows = [
        ('训练样本', '44 万帧', '百万～亿级'),
        ('被试数', '7 人', '千～万人'),
        ('冷启动', '< 14 天', '不可用'),
        ('模型大小', '< 30MB', 'GB 级'),
        ('推理硬件', 'CPU < 100ms', 'GPU 集群'),
        ('增量学习', '自监督，每天变准', '需重新训练'),
        ('个人基线', '14 天建立', '不支持'),
    ]

    table_y = 1.85
    row_h = 0.42
    col1_x = rx
    col2_x = rx + 1.35
    col3_x = rx + 2.6
    col1_w = 1.3
    col2_w = 1.2
    col3_w = 1.5

    # 表头
    put(s, col1_x, table_y, col1_w, 0.22, '维度', 7.5, RGBColor(0x88, 0x88, 0x88), spc=0.8)
    put(s, col2_x, table_y, col2_w, 0.22, '家安', 7.5, IKB_BRT, bold=True, spc=0.8)
    put(s, col3_x, table_y, col3_w, 0.22, '典型大模型', 7.5, RGBColor(0x88, 0x88, 0x88), spc=0.8)
    hline(s, rx, table_y + 0.28, rw, RGBColor(0x44, 0x44, 0x44), 0.5)

    for i, (dim, us, them) in enumerate(table_rows):
        ry = table_y + 0.38 + i * row_h
        # 交替底色
        if i % 2 == 0:
            rect(s, rx - 0.05, ry - 0.02, rw + 0.1, row_h - 0.04,
                 fill=RGBColor(0x18, 0x18, 0x18), line=None)
        put(s, col1_x, ry, col1_w, 0.3, dim, 8, RGBColor(0xBB, 0xBB, 0xBB))
        put(s, col2_x, ry, col2_w, 0.3, us, 8.5, WHITE, bold=True, latin=SANS)
        put(s, col3_x, ry, col3_w, 0.3, them, 8, RGBColor(0x77, 0x77, 0x77), latin=SANS)

    # 底部说明
    hline(s, rx, table_y + 0.38 + 7 * row_h + 0.08, rw, RGBColor(0x44, 0x44, 0x44), 0.5)
    put(s, rx, table_y + 0.38 + 7 * row_h + 0.18, rw, 1.0,
        '毫米波点云维度极低（10–50 点/帧 vs 视觉的'
        '数万像素），信息密度反而更高——每帧已经是'
        '物理空间的稀疏摘要。加上自监督对比学习，'
        '每个正常帧都是免费标签。\n\n'
        '不是"数据少所以模型弱"，'
        '是"数据结构简单所以不需要大数据"。',
        7.5, RGBColor(0x88, 0x88, 0x88), lh=1.3)

    # ── 底部 IKB 结论条 ──
    conclusion_y = 6.2
    rect(s, M, conclusion_y, 12.23, 0.65, fill=IKB)
    tf = box(s, M + 0.22, conclusion_y + 0.08, 11.8, 0.5, MSO_ANCHOR.MIDDLE)
    p = tf.paragraphs[0]; p.line_spacing = 1.15
    add_run(p, '世界模型路线成立。', 12, WHITE, bold=True)
    add_run(p, '"预测偏差 = 异常"机制在公开数据集上验证通过——系统从正常行为中学出'
             '"什么是正常"，识别未见过的危险。Encoder + Predictor < 30MB，'
             'ONNX 导出，CPU 推理 < 100ms——实时运行在云端，无需 GPU。',
             8.5, RGBColor(0xD5, 0xDD, 0xF2))

    # ── 诚实备注 ──
    put(s, M, 6.98, 12.23, 0.4,
        '数据集：mmWave-3DPCHM-1.0（TI IWR1443-ISK + Vayyar vBlu，实验室环境）'
        '· 留一法 7 折交叉验证取均值 · 以上精度为学术验证结果，非产品实测',
        7, RGBColor(0x66, 0x66, 0x66), lh=1.2)


def build():
    prs = Presentation()
    prs.slide_width = EMU_W
    prs.slide_height = EMU_H
    build_slide(prs)
    out = '世界模型_训练验证.pptx'
    prs.save(out)
    print(f'✅ 已生成: {out} (1 页)')


if __name__ == '__main__':
    build()
